"""Bounded, lossless tabular imports and spreadsheet-safe result exports."""

import csv
import io
import json
import math
import re
import zipfile
from datetime import date, datetime, time
from pathlib import PurePath

from openpyxl import Workbook, load_workbook

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 256
MAX_SHEETS = 20
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class FileValidationError(ValueError):
    """A file cannot be safely interpreted as an import table."""


def _headers(values):
    if not values or len(values) > MAX_COLUMNS:
        raise FileValidationError(f"Use between 1 and {MAX_COLUMNS} columns.")
    headers = [str(value) if value is not None else "" for value in values]
    if any(not header.strip() for header in headers):
        raise FileValidationError("Every column must have a nonempty header.")
    if len(set(headers)) != len(headers):
        raise FileValidationError("Column headers must be unique.")
    return headers


def _json_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _csv(content, delimiter):
    if len(delimiter) != 1 or delimiter in '\r\n\x00"':
        raise FileValidationError("Choose a single valid CSV delimiter.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileValidationError("CSV files must use UTF-8 encoding.") from exc
    if "\x00" in text:
        raise FileValidationError("CSV contains invalid binary data.")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        headers = _headers(next(reader, []))
        rows = []
        for values in reader:
            if not values:  # Ignore physical blank lines, retain rows containing empty cells.
                continue
            if len(rows) >= MAX_ROWS:
                raise FileValidationError(f"Files may contain at most {MAX_ROWS:,} data rows.")
            if len(values) != len(headers):
                raise FileValidationError(
                    f"CSV record {len(rows) + 2} has {len(values)} cells; expected {len(headers)}."
                )
            rows.append(dict(zip(headers, values, strict=True)))
    except csv.Error as exc:
        raise FileValidationError(f"Invalid CSV: {exc}") from exc
    return dict(headers=headers, rows=rows, sheet_names=[], sheet_name=None)


def _xlsx(content, sheet_name):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise FileValidationError("XLSX contains too many archive entries.")
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise FileValidationError("Expanded XLSX exceeds the 100 MiB limit.")
            names = {entry.filename for entry in entries}
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise FileValidationError("The file is not an XLSX workbook.")
            if any("vbaproject" in name.lower() for name in names):
                raise FileValidationError("Macro-enabled workbooks are not supported.")
            expanded = 0
            for entry in entries:
                tail = b""
                with archive.open(entry) as stream:
                    while chunk := stream.read(64 * 1024):
                        expanded += len(chunk)
                        if expanded > MAX_EXPANDED_BYTES:
                            raise FileValidationError("Expanded XLSX exceeds the 100 MiB limit.")
                        if entry.filename.endswith((".xml", ".rels")):
                            probe = tail + chunk
                            # OOXML does not require DTDs. Refuse them before XML parsing,
                            # even when optional parser hardening packages are absent.
                            if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
                                raise FileValidationError("XLSX XML declarations are unsafe.")
                            tail = probe[-16:]
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            names = workbook.sheetnames
            if len(names) > MAX_SHEETS:
                raise FileValidationError(f"Workbooks may contain at most {MAX_SHEETS} sheets.")
            selected = sheet_name if sheet_name is not None else names[0]
            if selected not in names:
                raise FileValidationError("The selected worksheet does not exist.")
            sheet = workbook[selected]
            # Do not trust dimensions stored by the workbook producer.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            if any(cell.data_type == "f" for cell in first):
                raise FileValidationError("Formula cells are not supported; paste values first.")
            headers = _headers([cell.value for cell in first])
            rows = []
            for cells in iterator:
                if len(rows) >= MAX_ROWS:
                    raise FileValidationError(f"Files may contain at most {MAX_ROWS:,} data rows.")
                if len(cells) > len(headers) and any(
                    cell.value is not None for cell in cells[len(headers) :]
                ):
                    raise FileValidationError(
                        "Worksheet contains values beyond the header columns."
                    )
                if any(cell.data_type == "f" for cell in cells):
                    raise FileValidationError(
                        "Formula cells are not supported; paste values first."
                    )
                values = [_json_value(cell.value) for cell in cells[: len(headers)]]
                values.extend([""] * (len(headers) - len(values)))
                rows.append(dict(zip(headers, values, strict=True)))
            return dict(headers=headers, rows=rows, sheet_names=names, sheet_name=selected)
        finally:
            workbook.close()
    except FileValidationError:
        raise
    except Exception as exc:
        raise FileValidationError("The XLSX workbook is damaged or unsupported.") from exc


def parse_file(
    filename: str, content: bytes, *, sheet_name: str | None = None, delimiter: str = ","
) -> dict:
    """Read UTF-8 CSV or one XLSX sheet; duplicates and input order are preserved."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise FileValidationError("File exceeds the 20 MiB upload limit.")
    if not content:
        raise FileValidationError("The file is empty.")
    extension = PurePath(filename).suffix.lower()
    if extension == ".csv":
        return _csv(content, delimiter)
    if extension == ".xlsx":
        return _xlsx(content, sheet_name)
    raise FileValidationError("Upload a .csv or .xlsx file.")


def _safe(value):
    """Escape formula triggers, including triggers hidden behind control/space prefixes."""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    probe = value.lstrip("".join(chr(number) for number in range(33)))
    if probe.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        value = "'" + value
    # XML cannot represent these controls; replace explicitly instead of failing silently.
    return _ILLEGAL_XML.sub("\ufffd", value)


def export_rows(headers: list[str], rows: list[dict], format: str) -> bytes:
    """Export a caller-ordered snapshot, keeping original and generated columns distinct."""
    if format not in {"csv", "xlsx"}:
        raise FileValidationError("Export format must be csv or xlsx.")
    used = set(headers)
    columns = [(header, "input_json", header) for header in headers]

    def append_column(base, source, key):
        name, number = base, 2
        while name in used:
            name, number = f"{base}.{number}", number + 1
        used.add(name)
        columns.append((name, source, key))

    keys = dict.fromkeys(key for row in rows for key in (row.get("output_json") or {}))
    for key in keys:
        append_column(f"result.{key}", "output_json", key)
    for key in ("row_number", "status", "error_code", "error_message"):
        append_column(key, None, key)

    def values():
        yield [_safe(column[0]) for column in columns]
        for row in rows:
            yield [
                _safe((row.get(source) or {}).get(key, "") if source else row.get(key, ""))
                for _, source, key in columns
            ]

    if format == "csv":
        output = io.StringIO(newline="")
        csv.writer(output, lineterminator="\r\n").writerows(values())
        return output.getvalue().encode("utf-8-sig")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Results")
    for record in values():
        if any(isinstance(value, str) and len(value) > 32767 for value in record):
            raise FileValidationError("A cell exceeds Excel's 32,767-character limit; export CSV.")
        sheet.append(record)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
