import csv
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.services import files
from app.services.files import FileValidationError, export_rows, parse_file


def workbook_bytes(records, second_sheet=False):
    book = Workbook()
    for row in records:
        book.active.append(row)
    if second_sheet:
        book.create_sheet("Other").append(["other"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_csv_preserves_duplicate_records_order_quoted_values_and_bom():
    data = '\ufeffid,note\r\n001,"a,b"\r\n001,"a,b"\r\n002,"two\nlines"\r\n'
    parsed = parse_file("input.CSV", data.encode())
    assert parsed["headers"] == ["id", "note"]
    assert parsed["rows"] == [
        {"id": "001", "note": "a,b"},
        {"id": "001", "note": "a,b"},
        {"id": "002", "note": "two\nlines"},
    ]
    assert parsed["sheet_name"] is None


def test_csv_custom_delimiter_and_empty_values():
    parsed = parse_file("a.csv", b"id;name\n;\n\n", delimiter=";")
    assert parsed["rows"] == [{"id": "", "name": ""}]


@pytest.mark.parametrize(
    "content",
    [b"id,id\n1,2", b"id,\n1,2", b"id,name\n1", b'id\n"unclosed', b"id\n\xff", b"id\n\x00"],
)
def test_csv_rejects_invalid_structure(content):
    with pytest.raises(FileValidationError):
        parse_file("a.csv", content)


@pytest.mark.parametrize("delimiter", ["", "aa", "\n", '"'])
def test_invalid_delimiter(delimiter):
    with pytest.raises(FileValidationError):
        parse_file("a.csv", b"id\n1", delimiter=delimiter)


def test_file_and_row_limits(monkeypatch):
    monkeypatch.setattr(files, "MAX_UPLOAD_BYTES", 5)
    with pytest.raises(FileValidationError, match="upload limit"):
        parse_file("a.csv", b"id\n123")
    monkeypatch.setattr(files, "MAX_UPLOAD_BYTES", 1_000_000)
    monkeypatch.setattr(files, "MAX_ROWS", 2)
    assert len(parse_file("a.csv", b"id\n1\n2")["rows"]) == 2
    with pytest.raises(FileValidationError, match="data rows"):
        parse_file("a.csv", b"id\n1\n2\n3")
    with pytest.raises(FileValidationError, match="data rows"):
        parse_file("a.xlsx", workbook_bytes([["id"], [1], [2], [3]]))


def test_xlsx_sheet_selection_and_values():
    content = workbook_bytes([["id", "quantity"], ["001", 2], ["001", 2]], True)
    parsed = parse_file("a.xlsx", content)
    assert parsed["sheet_names"] == ["Sheet", "Other"]
    assert parsed["sheet_name"] == "Sheet"
    assert parsed["rows"] == [{"id": "001", "quantity": 2}] * 2
    assert parse_file("a.xlsx", content, sheet_name="Other")["headers"] == ["other"]
    with pytest.raises(FileValidationError, match="does not exist"):
        parse_file("a.xlsx", content, sheet_name="Missing")


def test_xlsx_formula_rejected_even_without_cached_values():
    with pytest.raises(FileValidationError, match="Formula"):
        parse_file("a.xlsx", workbook_bytes([["id"], ["=1+2"]]))


def test_corrupt_and_wrong_file_types():
    for name, content in [
        ("a.xlsx", b"not a zip"),
        ("a.xlsm", b"data"),
        ("a.csv", b""),
        ("a.xlsx", b"id\n1"),
    ]:
        with pytest.raises(FileValidationError):
            parse_file(name, content)


def test_expansion_and_macro_checks(monkeypatch):
    content = workbook_bytes([["id"], [1]])
    monkeypatch.setattr(files, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(FileValidationError, match="Expanded XLSX"):
        parse_file("a.xlsx", content)
    monkeypatch.setattr(files, "MAX_EXPANDED_BYTES", 100 * 1024 * 1024)
    output = io.BytesIO(content)
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(FileValidationError, match="Macro"):
        parse_file("a.xlsx", output.getvalue())


def test_xlsx_rejects_xml_entity_declarations():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("xl/workbook.xml", b'<!DOCTYPE x [<!ENTITY y "z">]><x/>')
    with pytest.raises(FileValidationError, match="unsafe"):
        parse_file("a.xlsx", output.getvalue())


def test_three_thousand_records_retained_exactly_once():
    content = "id,quantity\n" + "".join(f"{index % 10},{index}\n" for index in range(3000))
    parsed = parse_file("input.csv", content.encode())
    rows = [
        {"input_json": row, "row_number": index, "status": "pending"}
        for index, row in enumerate(parsed["rows"], 1)
    ]
    exported = list(
        csv.DictReader(io.StringIO(export_rows(parsed["headers"], rows, "csv").decode("utf-8-sig")))
    )
    assert len(exported) == 3000
    assert [record["quantity"] for record in exported] == [str(index) for index in range(3000)]


def test_export_preserves_inputs_collisions_and_duplicates():
    headers = ["id", "status", "result.price", "status.2"]
    row = {
        "input_json": {
            "id": "001",
            "status": "input",
            "result.price": "original",
            "status.2": "also original",
        },
        "output_json": {"price": 12},
        "status": "succeeded",
        "row_number": 1,
    }
    result = list(
        csv.reader(io.StringIO(export_rows(headers, [row, row], "csv").decode("utf-8-sig")))
    )
    assert result[0] == headers + [
        "result.price.2",
        "row_number",
        "status.3",
        "error_code",
        "error_message",
    ]
    assert result[1] == result[2]
    assert result[1][:4] == ["001", "input", "original", "also original"]
    assert result[1][4:7] == ["12", "1", "succeeded"]


@pytest.mark.parametrize(
    "value", ["=1+1", "+cmd", "-cmd", "@SUM(A1)", "\t=1", "\r+1", "\n@1", "  =1", "\x01=1"]
)
@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_formula_injection_protection(value, format):
    data = export_rows(
        [value],
        [{"input_json": {value: value}, "output_json": {"danger": value}, "error_message": value}],
        format,
    )
    if format == "csv":
        result = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        book = load_workbook(io.BytesIO(data), data_only=False)
        assert all(cell.data_type != "f" for row in book.active for cell in row)
        result = list(book.active.values)
        book.close()
    assert result[0][0].startswith("'")
    assert result[1][0].startswith("'")
    assert result[1][1].startswith("'")
    assert result[1][-1].startswith("'")


def test_xlsx_export_numeric_values_and_partial_status():
    data = export_rows(
        ["id"],
        [
            {
                "input_json": {"id": 1},
                "output_json": {"price": -2},
                "status": "pending",
                "row_number": 1,
            }
        ],
        "xlsx",
    )
    book = load_workbook(io.BytesIO(data))
    records = list(book.active.values)
    assert records[1][:4] == (1, -2, 1, "pending")
    book.close()


def test_export_unsupported_format():
    with pytest.raises(FileValidationError):
        export_rows([], [], "pdf")
